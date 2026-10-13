from __future__ import annotations

import csv
import gzip
import hashlib
import json
import math
from collections import Counter
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / 'data/03_Ecoli_strength87'
TARGET = ROOT / 'data/04_Bsub_strength61'
IDENTITY = dict(dataset_id='course_ecoli87_strength', data_version='ecoli87_tss_verified_v1',
                split_id='ecoli87_inherited_20260928_v1')


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def digest_ids(ids):
    return hashlib.sha256(''.join(x+'\n' for x in sorted(ids)).encode()).hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False)+'\n', encoding='utf-8')


def read_table(path):
    with Path(path).open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f, delimiter='\t' if Path(path).suffix=='.tsv' else ','))


def write_table(path, rows, columns=None):
    rows=list(rows); path=Path(path); path.parent.mkdir(parents=True,exist_ok=True)
    with path.open('w',encoding='utf-8',newline='') as f:
        writer=csv.DictWriter(f, fieldnames=columns or list(rows[0]),
                              delimiter='\t' if path.suffix=='.tsv' else ',',lineterminator='\n')
        writer.writeheader(); writer.writerows(rows)


def require(condition, message):
    if not condition: raise ValueError(message)


def sequence_check(sequence, length):
    require(len(sequence)==length and set(sequence)<=set('ACGT'),f'Expected {length} uppercase ACGT bases')


def rc(sequence):
    return sequence.translate(str.maketrans('ACGT','TGCA'))[::-1]


def circular(genome,start,length):
    return ''.join(genome[(start+i)%len(genome)] for i in range(length))


def coordinate(position, tss=67):
    """Zero-based sequence index -> biological position, skipping position zero."""
    return position-tss if position<tss else position-tss+1


def records():
    with gzip.open(BASE/'raw/verified_context.jsonl.gz','rt') as f:
        return [json.loads(line) for line in f]


def reference_hits(genome,sequences):
    """Recompute all exact matches on both strands, including circular-origin matches."""
    from collections import defaultdict
    seeds=defaultdict(list)
    for i,sequence in enumerate(sequences):
        sequence=str(sequence).upper()
        for strand,pattern in [('+',sequence),('-',rc(sequence))]:
            seeds[pattern[:16]].append((i,strand,pattern))
    hits=[[] for _ in sequences]; search=genome+genome[:49]
    for start in range(len(genome)):
        for i,strand,pattern in seeds.get(search[start:start+16],()):
            if search[start:start+50]==pattern: hits[i].append((start+1,strand))
    return hits


def source_check(full=True):
    profile=read_json(BASE/'source_manifest.json')
    for name,expected in profile['sha256'].items():
        require(sha(ROOT/name)==expected,'Source hash mismatch: '+name)
    recs=records()
    original=np.load(ROOT/'data/raw_ecoli50/promoter.npy',allow_pickle=False)
    labels=np.load(ROOT/'data/raw_ecoli50/gene_expression.npy',allow_pickle=False)
    oldsplit={r['sample_id']:r['split'] for r in read_table(ROOT/'data/01_Ecoli_strength/split_manifest.tsv')}
    with gzip.open(BASE/'raw/U00096.2.fasta.gz','rt') as f:
        genome=''.join(line.strip() for line in f if not line.startswith('>')).upper()
    require(len(genome)==4639675,'Unexpected historical reference length')
    tss_sequences={}
    if full:
        hits=reference_hits(genome,original)
        for record in recs:
            require(hits[record['source_index_0based']]==[(record['match_start_1based'],record['strand'])],
                    'Source sequence is not uniquely mapped to declared position')
        from openpyxl import load_workbook
        workbook=load_workbook(BASE/'raw/thomason_dataset1.xlsx',read_only=True,data_only=True)
        for values in workbook['TSS Map MasterTable'].iter_rows(min_row=4,values_only=True):
            if isinstance(values[0],(int,float)) and values[1] in ('+','-') and isinstance(values[23],str):
                tss_sequences.setdefault((int(values[0]),values[1]),set()).add(values[23].upper())
        workbook.close()
    seen=set()
    for r in recs:
        i=r['source_index_0based']; seq=r['sequence_87bp']; start=r['match_start_1based']-1
        require(i not in seen,'Duplicate original row'); seen.add(i)
        sequence_check(seq,87)
        require(seq[17:67]==str(original[i]).upper()==r['source_sequence_50bp'],'Original 50 bp mismatch')
        require(str(labels[i])==r['expression_original_text'] and float(labels[i])==r['expression'], 'Original label mismatch')
        expected=circular(genome,start-17,87) if r['strand']=='+' else rc(circular(genome,start-20,87))
        require(seq==expected,'Reference context mismatch')
        tss=(start+50)%len(genome)+1 if r['strand']=='+' else (start-1)%len(genome)+1
        require(r['tss_1based']==tss and r['tss_index_0based_in_87bp']==67,'TSS coordinate mismatch')
        if full: require(seq[17:68] in tss_sequences.get((tss,r['strand']),set()),'Original experimental TSS table mismatch')
        require(f'ecoli50_r{i+1:06d}' in oldsplit,'Missing parent split')
    with gzip.open(BASE/'raw/excluded_ambiguous.jsonl.gz','rt') as f:
        excluded=[json.loads(line)['source_index_0based'] for line in f]
    require(len(recs)==11809 and len(excluded)==75 and seen.isdisjoint(excluded)
            and seen|set(excluded)==set(range(len(original))),'Exclusions do not reconcile with original dataset')
    if full: require(all(len(hits[i])>1 for i in excluded),'Excluded row is not an ambiguous genomic match')
    return dict(status='passed',n_retained=len(recs),n_excluded=len(excluded),reference='U00096.2',
                original_labels_unchanged=True,original_split_inherited=True,
                original_tss_table_checked=full)


def transformed(rows):
    train=[r for r in rows if r['split']=='train']
    y=np.log10([float(r['strength']) for r in train])
    return dict(**IDENTITY,kind='log10_minmax',fit_subset='train',clip=False,
                train_ids_sha256=digest_ids(r['sample_id'] for r in train),n_fit=len(train),
                min_log10=float(y.min()),max_log10=float(y.max()))


def build(output=BASE, verify=True):
    if verify: source_check()
    output=Path(output)
    require(not (output/'samples.tsv').exists(),'Dataset destination already exists')
    parent={r['sample_id']:r['split'] for r in read_table(ROOT/'data/01_Ecoli_strength/split_manifest.tsv')}
    rows=[]
    for r in records():
        index=r['source_index_0based']; old_id=f'ecoli50_r{index+1:06d}'
        rows.append(dict(**IDENTITY,sample_id=f'ecoli87_r{index+1:06d}',source_row=index+1,
                         parent_sample_id=old_id,sequence=r['sequence_87bp'],sequence_length=87,
                         strength=r['expression'],target_log10=math.log10(r['expression']),
                         split=parent[old_id],tss_index_0based=67,tss_1based=r['tss_1based'],
                         reference=r['reference'],strand=r['strand'],annotation_status='experimental_tss_verified'))
    write_table(output/'samples.tsv',rows)
    write_table(output/'split_manifest.tsv',[{k:r[k] for k in ['dataset_id','data_version','split_id','sample_id','parent_sample_id','split']} for r in rows])
    write_json(output/'label_transform.json',transformed(rows))
    write_json(output/'data_manifest.json',dict(**IDENTITY,n_input=11884,n_retained=11809,n_excluded_ambiguous=75,
               sequence_length=87,tss_index_0based=67,original_50bp_slice=[17,67],
               split_policy='inherit original sample assignment; no redraw',counts=dict(Counter(r['split'] for r in rows)),
               source_manifest_sha256=sha(BASE/'source_manifest.json'),
               annotation_policy='TSS independently verified; boxes remain model inferred or canonical assumptions',
               label_policy='unchanged original labels; log10 minmax fitted on retained training rows only'))
    return rows


def load(cfg=None):
    cfg=cfg or read_json(ROOT/'project_config.json')
    rows=read_table(ROOT/cfg['samples'])
    splits=read_table(ROOT/cfg['splits'])
    transform=read_json(ROOT/cfg['label_transform'])
    require(cfg['schema_version']=='3.0.0' and cfg['test_metrics_enabled'] is False,'Invalid context workflow configuration')
    require(cfg['sequence_length']==87 and cfg['tss_index_0based']==67,'Verified context uses length 87, TSS index 67')
    ids=[r['sample_id'] for r in rows]
    require(len(ids)==cfg['expected_samples'] and len(set(ids))==len(ids),'Incorrect sample IDs/count')
    require(len({r['sequence'] for r in rows})==len(rows),'Duplicate main sequences')
    require(len(splits)==len(rows) and {r['sample_id'] for r in splits}==set(ids),'Incomplete split manifest')
    require(len({r['sample_id'] for r in splits})==len(splits),'Duplicate split IDs')
    splitmap={r['sample_id']:r for r in splits}
    for r in rows:
        sequence_check(r['sequence'],cfg['sequence_length'])
        require(int(r['sequence_length'])==87 and int(r['tss_index_0based'])==67,'Invalid row coordinates')
        require(r['split'] in ('train','val','test') and splitmap[r['sample_id']]['split']==r['split'],'Split mismatch')
        for k in IDENTITY: require(r[k]==cfg[k]==transform[k]==splitmap[r['sample_id']][k],'Data identity mismatch: '+k)
        strength=float(r['strength']); require(math.isfinite(strength) and strength>0,'Invalid strength')
        require(math.isclose(float(r['target_log10']),math.log10(strength),abs_tol=1e-12),'Incorrect log10 label')
    require(transform==transformed(rows),'Transform is not fitted on current training rows')
    return rows,transform


def target_build():
    with (TARGET/'raw/promoter_strength.csv').open(encoding='utf-8-sig',newline='') as f:
        records_in=list(csv.reader(f))
    rows=[]
    for i,record in enumerate(records_in,1):
        require(len(record)==2,'Target CSV requires two headerless columns')
        sequence=record[0].strip().upper(); strength=float(record[1]); sequence_check(sequence,61)
        require(math.isfinite(strength) and strength>0,'Invalid target label')
        rows.append(dict(sample_id=f'bsub61_r{i:06d}',sequence=sequence,sequence_length=61,
                         strength=strength,target_log10=math.log10(strength),
                         sequence_group=hashlib.sha256(sequence.encode()).hexdigest()[:20]))
    require(len(rows)==591 and len({r['sequence_group'] for r in rows})==580,'Unexpected teacher CSV contents')
    write_table(TARGET/'samples.tsv',rows)
    write_json(TARGET/'data_manifest.json',dict(dataset_id='teacher_bsub61_strength',n_records=591,n_unique_sequences=580,
               source_sha256=sha(TARGET/'raw/promoter_strength.csv'),sequence_length=61,
               units='unknown',tss='unknown',conditions='unknown',duplicate_policy='retain measurements; split by exact sequence group'))
    return rows


def validate(cfg=None,full=True):
    rows,transform=load(cfg)
    proof=source_check(full)
    expected=records()
    for row,raw in zip(rows,expected):
        require(int(row['source_row'])==raw['source_index_0based']+1 and row['sequence']==raw['sequence_87bp']
                and float(row['strength'])==raw['expression'],'Canonical dataset differs from verified source')
        oldid=f"ecoli50_r{int(row['source_row']):06d}"
        require(row['sample_id']==f"ecoli87_r{int(row['source_row']):06d}" and row['parent_sample_id']==oldid,'Lineage ID mismatch')
        require(row['reference']==raw['reference'] and row['strand']==raw['strand']
                and int(row['tss_1based'])==raw['tss_1based'],'Genomic coordinate annotation mismatch')
    target=read_json(TARGET/'data_manifest.json')
    require(sha(TARGET/'raw/promoter_strength.csv')==target['source_sha256'],'Target source changed')
    with (TARGET/'raw/promoter_strength.csv').open(encoding='utf-8-sig',newline='') as f: original_target=list(csv.reader(f))
    current_target=read_table(TARGET/'samples.tsv')
    require(len(current_target)==len(original_target)==591,'Target record count changed')
    for i,(r,source) in enumerate(zip(current_target,original_target),1):
        require(r['sample_id']==f'bsub61_r{i:06d}' and r['sequence']==source[0].strip().upper()
                and float(r['strength'])==float(source[1]),'Target source record changed')
        require(r['sequence_group']==hashlib.sha256(r['sequence'].encode()).hexdigest()[:20], 'Target sequence group mismatch')
    parent={r['sample_id']:r['split'] for r in read_table(ROOT/'data/01_Ecoli_strength/split_manifest.tsv')}
    require(all(parent[r['parent_sample_id']]==r['split'] for r in rows),'Inherited split changed')
    return dict(**proof,counts=dict(Counter(r['split'] for r in rows)),transform_train_only=True)
