"""Typeset the structured course report as an editable Word document."""
from __future__ import annotations

import argparse
import json
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
CJK='Source Han Sans CN'
BLACK=RGBColor(0,0,0)


def font(run,size=11,bold=False):
    run.font.name='Arial';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=BLACK
    fonts=run._element.get_or_add_rPr().get_or_add_rFonts()
    for key,value in [('ascii','Arial'),('hAnsi','Arial'),('eastAsia',CJK),('cs','Arial')]:fonts.set(qn('w:'+key),value)


def build(content_path,output):
    data=json.loads(content_path.read_text());doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.72);sec.left_margin=sec.right_margin=Inches(.8)
    sec.footer_distance=Inches(.28)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        style=doc.styles[name];style.font.name='Arial';style.font.size=Pt(11);style.font.color.rgb=BLACK
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),CJK)
        style.paragraph_format.line_spacing=Pt(16.3);style.paragraph_format.space_after=Pt(7)
        for border in list(style._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for name,size in [('Title',24),('Heading 1',16),('Heading 2',12)]:
        style=doc.styles[name];style.font.size=Pt(size);style.font.bold=True
        style.paragraph_format.line_spacing=Pt(size*1.3);style.paragraph_format.space_after=Pt(10)
        style.paragraph_format.keep_with_next=True
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(8)
    doc.styles['Caption'].font.size=Pt(9.5)
    doc.styles['Caption'].paragraph_format.line_spacing=Pt(13)
    def para(text,size=11,style=None,bold=False):
        p=doc.add_paragraph(style=style);p.paragraph_format.widow_control=True
        grid=OxmlElement('w:snapToGrid');grid.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(grid)
        font(p.add_run(text),size,bold);return p
    title=doc.add_paragraph(data['title'],style='Title')
    for run in title.runs:font(run,24,True)
    para(data['subtitle'],12,bold=True)
    para('项目组  '+'、'.join(data['members']),10)
    para(data['date'],10)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font(footer.add_run('PR02-01 研究报告  '),9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

    def table(block):
        caption=para(f"表{block['number']} {block['caption']}",10,style='Caption',bold=True)
        caption.paragraph_format.keep_with_next=True
        n=len(block['headers']);widths=block.get('widths') or [6.7/n]*n
        t=doc.add_table(rows=1,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        for col,width in zip(t.columns,widths):col.width=Inches(width)
        props=t._tbl.tblPr;tw=props.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(round(sum(widths)*1440)))
        borders=OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        props.append(borders);margins=OxmlElement('w:tblCellMar')
        for side,value in [('top',85),('bottom',85),('left',90),('right',90)]:
            e=OxmlElement('w:'+side);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');margins.append(e)
        props.append(margins)
        for i,values in enumerate([block['headers']]+block['rows']):
            row=t.rows[0] if i==0 else t.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for j,(cell,value) in enumerate(zip(row.cells,values)):
                cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'DDE7F0' if i==0 else 'FFFFFF');cell._tc.get_or_add_tcPr().append(shading)
                p=cell.paragraphs[0];p.paragraph_format.line_spacing=Pt(13);p.paragraph_format.space_after=Pt(0)
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 or n>=4 and j>0 else WD_ALIGN_PARAGRAPH.LEFT
                font(p.add_run(str(value)),9.6,i==0)
        after=para('',size=3);after.paragraph_format.line_spacing=Pt(3);after.paragraph_format.space_after=Pt(2)

    for i,section in enumerate(data['sections']):
        p=doc.add_paragraph(section['title'],style='Heading 1')
        if i and section.get('new_page',True):p.paragraph_format.page_break_before=True
        for block in section['blocks']:
            kind=block['type']
            if kind=='paragraph':para(block['text'])
            elif kind=='subheading':para(block['text'],12,'Heading 2',True)
            elif kind=='table':table(block)
            elif kind=='figure':
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_with_next=True
                image=content_path.parent/block['path'];p.add_run().add_picture(str(image),width=Inches(6.7))
                props=p.runs[0]._r.xpath('.//wp:docPr')
                if props:props[0].set('descr',block['title'])
                para(f"图{block['number']} {block['title']}。{block['caption']}",9.5,'Caption')
            elif kind=='reference':
                p=para(f"[{block['number']}] {block['text']}",10)
                p.paragraph_format.keep_with_next=True
                p=para(block['url'],9.5);p.paragraph_format.space_after=Pt(12)
            elif kind=='code':
                for line in block['text'].splitlines():
                    p=para(line,9);p.paragraph_format.line_spacing=Pt(13);p.paragraph_format.space_after=Pt(2)
                    p.runs[0].font.name='Courier New'
    cp=doc.core_properties;cp.title=data['title'];cp.subject='PR02-01 启动子强度预测与关键位点归因'
    cp.author='PR02-01 项目组';cp.last_modified_by='';cp.comments='';cp.keywords='启动子强度 CNN 热力学 归因'
    cp.created=datetime(2026,10,14,tzinfo=timezone.utc);cp.modified=cp.created;cp.revision=1
    output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)
    # Remove template and editing-session metadata while preserving semantic document content.
    with zipfile.ZipFile(output) as archive:parts={n:archive.read(n) for n in archive.namelist()}
    app=ElementTree.fromstring(parts['docProps/app.xml']);namespace='{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}'
    for name in ['Application','AppVersion','Company','Manager','Template']:
        node=app.find(namespace+name)
        if node is not None:node.text=''
    parts['docProps/app.xml']=ElementTree.tostring(app,encoding='utf-8',xml_declaration=True)
    with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as archive:
        for name,body in parts.items():archive.writestr(name,body)
    print(output)


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--content',type=Path,default=ROOT/'reports/final/report_content.json')
    parser.add_argument('--output',type=Path,default=ROOT/'reports/final/PR02-01_REPORT.docx')
    args=parser.parse_args();build(args.content,args.output)
