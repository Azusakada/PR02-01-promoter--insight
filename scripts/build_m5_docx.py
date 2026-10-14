"""Render M5_REPORT_CONTENT.json into an editable report (python-docx required)."""
from pathlib import Path
import argparse
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

FONT='Source Han Sans CN'


def set_font(run,size=11,bold=False,color='000000'):
    run.font.name='Arial';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
    rf=run._element.get_or_add_rPr().get_or_add_rFonts()
    for name,value in [('ascii','Arial'),('hAnsi','Arial'),('eastAsia',FONT),('cs','Arial')]:rf.set(qn('w:'+name),value)


def build(repo,output):
    data=json.loads((repo/'reports/M5_REPORT_CONTENT.json').read_text())
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.68)
    sec.left_margin=sec.right_margin=Inches(.8);sec.footer_distance=Inches(.27)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        style=doc.styles[name];style.font.name='Arial';style.font.size=Pt(11);style.font.color.rgb=RGBColor(0,0,0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONT)
        style.paragraph_format.line_spacing=Pt(15.3);style.paragraph_format.space_after=Pt(6)
        for border in list(style._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for name,size in [('Title',25),('Heading 1',17),('Heading 2',12)]:
        s=doc.styles[name];s.font.size=Pt(size);s.font.bold=True;s.paragraph_format.keep_with_next=True
        s.paragraph_format.line_spacing=Pt(size*1.3);s.paragraph_format.space_before=Pt(3);s.paragraph_format.space_after=Pt(10)
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    set_font(foot.add_run('PR02-01 M5 最终报告    '),9,color='555555')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    def paragraph(text,size=11,color='000000',bold=False):
        p=doc.add_paragraph();p.paragraph_format.widow_control=True
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(snap)
        set_font(p.add_run(text),size,bold,color);return p
    def table(source):
        headers=source['headers'];n=len(headers)
        widths={2:[1.6,5.3],3:[2.2,1.55,3.15],4:[1.85,1.2,1.35,2.5],5:[1.25,1.5,1.5,1.25,1.4]}[n]
        t=doc.add_table(rows=1,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        for col,width in zip(t.columns,widths):col.width=Inches(width)
        props=t._tbl.tblPr;tw=props.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),'9936')
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'5');el.set(qn('w:color'),'D9D9D9');borders.append(el)
        props.append(borders)
        margins=OxmlElement('w:tblCellMar')
        for side,value in [('top',75),('bottom',75),('left',85),('right',85)]:
            el=OxmlElement('w:'+side);el.set(qn('w:w'),str(value));el.set(qn('w:type'),'dxa');margins.append(el)
        props.append(margins)
        for i,values in enumerate([headers]+source['rows']):
            row=t.rows[0] if i==0 else t.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for j,(cell,value) in enumerate(zip(row.cells,values)):
                cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DDE7F0' if i==0 else 'F6F8FA' if i%2==0 else 'FFFFFF');cell._tc.get_or_add_tcPr().append(sh)
                p=cell.paragraphs[0];p.paragraph_format.line_spacing=Pt(14);p.paragraph_format.space_after=Pt(0)
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 or n>=4 and j>0 else WD_ALIGN_PARAGRAPH.LEFT
                set_font(p.add_run(str(value)),10.1,i==0)
        p=doc.add_paragraph();p.paragraph_format.line_spacing=Pt(4);p.paragraph_format.space_after=Pt(2)
    title=doc.add_paragraph(data['title'],style='Title')
    for run in title.runs:set_font(run,25,True)
    paragraph(data['subtitle'],13,bold=True);paragraph(data['date']+'    项目组：'+'、'.join(data['members']),10,color='444444')
    for index,section in enumerate(data['sections']):
        heading=doc.add_paragraph(section['title'],style='Heading 1')
        if index:heading.paragraph_format.page_break_before=True
        for text in section['paragraphs']:paragraph(text)
        if section['table']:table(section['table'])
        if section['image']:
            width=6.1 if 'mutation' in section['image'] else 6.7
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4)
            p.paragraph_format.line_spacing=1
            p.add_run().add_picture(str(repo/section['image']),width=Inches(width))
        if section['sources']:
            names='、'.join(section['sources']);p=paragraph('源表：'+names,9,color='555555');p.paragraph_format.line_spacing=Pt(11.5)
    doc.core_properties.title=data['title'];doc.core_properties.subject='M4四项必做、选做5/6与M5最终交付'
    doc.core_properties.author='PR02-01 项目组'
    output.parent.mkdir(parents=True,exist_ok=True);doc.save(output);print(output)


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--repo',type=Path,default=Path(__file__).resolve().parents[1]);parser.add_argument('--output',type=Path)
    args=parser.parse_args();build(args.repo,args.output or args.repo/'reports/M5_FINAL_REPORT.docx')
